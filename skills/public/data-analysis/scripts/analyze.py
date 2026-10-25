"""
Data Analysis Script using DuckDB.

Analyzes Excel (.xlsx/.xls) and CSV files using DuckDB's in-process SQL engine.
Supports schema inspection, SQL queries, statistical summaries, and result export.
"""

import argparse
from contextlib import contextmanager
import hashlib
import json
import logging
import os
import re
import shutil
import tempfile
import sys
import uuid
from collections import Counter
from pathlib import Path

logging.basicConfig(level=logging.INFO, format="%(message)s")
logger = logging.getLogger(__name__)

try:
    import duckdb
except ImportError:
    raise SystemExit(
        "Missing duckdb. Prepare this Python environment with: python -m pip install duckdb openpyxl"
    ) from None

try:
    import openpyxl  # noqa: F401
except ImportError:
    raise SystemExit(
        "Missing openpyxl. Prepare this Python environment with: python -m pip install duckdb openpyxl"
    ) from None

# Cache directory for persistent DuckDB databases
CACHE_DIR = os.path.join(tempfile.gettempdir(), ".data-analysis-cache")


def compute_files_hash(files: list[str]) -> str:
    """Hash the ordered names and contents that determine the loaded table map."""
    # A new namespace also prevents reuse of maps built by the old allocator.
    hasher = hashlib.sha256(b"data-analysis-inputs-v2\0")
    for file_path in files:
        name = os.fsencode(os.path.basename(file_path))
        hasher.update(len(name).to_bytes(8, "big"))
        hasher.update(name)
        file_hasher = hashlib.sha256()
        try:
            with open(file_path, "rb") as f:
                while chunk := f.read(8192):
                    file_hasher.update(chunk)
        except OSError:
            # Keep unreadable paths distinct, including POSIX surrogate escapes.
            hasher.update(b"missing\0")
            file_hasher = hashlib.sha256(os.fsencode(file_path))
        else:
            hasher.update(b"readable\0")
        # Fixed-width digests preserve file boundaries, even for empty files.
        hasher.update(file_hasher.digest())
    return hasher.hexdigest()


def _open_cached_generation(root: Path, files_hash: str):
    """Open one complete generation; invalid cache entries are ordinary misses."""
    con = None
    try:
        manifest = json.loads((root / f"{files_hash}.json").read_text(encoding="utf-8"))
        if (
            not isinstance(manifest, dict)
            or type(manifest.get("format")) is not int
            or manifest["format"] != 1
        ):
            return None
        generation = manifest.get("generation")
        table_map = manifest.get("table_map")
        if not isinstance(generation, str) or not re.fullmatch(
            rf"{files_hash}\.[0-9a-f]{{32}}", generation
        ):
            return None
        if (
            not isinstance(table_map, dict)
            or not table_map
            or not all(
                isinstance(k, str) and isinstance(v, str) and k and v
                for k, v in table_map.items()
            )
        ):
            return None
        aliases = {name.casefold() for name in table_map}
        targets = {name.casefold() for name in table_map.values()}
        if len(aliases) != len(table_map) or len(targets) != len(table_map):
            return None
        if any(
            alias.casefold() in targets and alias.casefold() != target.casefold()
            for alias, target in table_map.items()
        ):
            return None
        directory = root / generation
        database = directory / "data.duckdb"
        if directory.is_symlink() or database.is_symlink():
            return None
        con = duckdb.connect(str(database), read_only=True)
        physical_tables = {
            row[0].casefold()
            for row in con.execute(
                "SELECT table_name FROM information_schema.tables "
                "WHERE table_schema = 'main' AND table_type = 'BASE TABLE' "
                "AND table_catalog = current_database()"
            ).fetchall()
        }
        if targets != physical_tables:
            raise ValueError("Cached table map does not match its database")
        for table in table_map.values():
            quoted = '"' + table.replace('"', '""') + '"'
            con.execute(f"SELECT * FROM main.{quoted} LIMIT 0")
        return con, table_map
    except (OSError, ValueError, duckdb.Error):
        if con is not None:
            con.close()
        logger.debug("Ignoring unreadable analysis cache", exc_info=True)
        return None
    except BaseException:
        if con is not None:
            con.close()
        raise


def _publish_cache_manifest(
    root: Path, files_hash: str, generation: str, table_map: dict[str, str]
) -> None:
    """Publish the database identity and map together, after its writer closed."""
    fd, temporary = tempfile.mkstemp(prefix=".manifest-", suffix=".tmp", dir=root)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(
                {"format": 1, "generation": generation, "table_map": table_map},
                handle,
                ensure_ascii=False,
            )
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, root / f"{files_hash}.json")
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


@contextmanager
def open_analysis_cache(files: list[str]):
    """Own a read-only analysis connection backed by an immutable generation."""
    files_hash = compute_files_hash(files)
    root = Path(CACHE_DIR) / "generations-v1"
    root.mkdir(parents=True, exist_ok=True)
    cached = _open_cached_generation(root, files_hash)
    if cached is not None:
        con, table_map = cached
        logger.info("Cache hit! Using complete cached generation")
        try:
            yield con, table_map
        finally:
            con.close()
        return

    generation = f"{files_hash}.{uuid.uuid4().hex}"
    directory = root / generation
    directory.mkdir(mode=0o700)
    database = directory / "data.duckdb"
    con = None
    publication_started = False
    try:
        logger.info("Loading files into a private cache generation...")
        con = duckdb.connect(str(database))
        table_map = load_files(con, files)
        if not table_map:
            logger.error("No tables were loaded. Check file paths and formats.")
            raise SystemExit(1)
        con.close()
        con = None
        # Never expose a database still owned by a writer, including the first
        # invocation's analysis action. Other processes may already be readers.
        con = duckdb.connect(str(database), read_only=True)
        # Retain complete generations once publication may have happened, even
        # if interrupted immediately after replace. A reader may have its ID.
        publication_started = True
        _publish_cache_manifest(root, files_hash, generation, table_map)
        logger.info("Loaded %s table(s): %s", len(table_map), ", ".join(table_map))
        yield con, table_map
    finally:
        if con is not None:
            con.close()
        if not publication_started:
            shutil.rmtree(directory)


def sanitize_table_name(name: str) -> str:
    """Sanitize a sheet/file name into a valid SQL table name."""
    sanitized = re.sub(r"[^\w]", "_", name)
    if sanitized and sanitized[0].isdigit():
        sanitized = f"t_{sanitized}"
    return sanitized


def _quote_identifier(name: str) -> str:
    """Quote a SQL identifier, including embedded double quotes."""
    return '"' + name.replace('"', '""') + '"'


def _allocate_table_names(name: str, table_map: dict[str, str]) -> tuple[str, str]:
    """Keep lookup aliases and SQL names unique, including case-only collisions."""
    reserved = {n.casefold() for n in (*table_map.keys(), *table_map.values())}
    base = sanitize_table_name(name)
    table_name = base
    counter = 1
    while table_name.casefold() in reserved:
        table_name = f"{base}_{counter}"
        counter += 1
    # An original name must never redirect an earlier canonical SQL name.
    alias = table_name if name.casefold() in reserved else name
    return alias, table_name


def load_files(con: duckdb.DuckDBPyConnection, files: list[str]) -> dict[str, str]:
    """
    Load Excel/CSV files into DuckDB tables.

    Returns unique lookup aliases mapped to their canonical SQL table names.
    """
    con.execute("INSTALL spatial; LOAD spatial;")
    table_map: dict[str, str] = {}

    for file_path in files:
        if not os.path.exists(file_path):
            logger.error(f"File not found: {file_path}")
            continue

        ext = os.path.splitext(file_path)[1].lower()

        if ext in (".xlsx", ".xls"):
            _load_excel(con, file_path, table_map)
        elif ext == ".csv":
            _load_csv(con, file_path, table_map)
        else:
            logger.warning(f"Unsupported file format: {ext} ({file_path})")

    return table_map


def _load_excel(
    con: duckdb.DuckDBPyConnection, file_path: str, table_map: dict[str, str]
) -> None:
    """Load all sheets from an Excel file into DuckDB tables."""
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    for sheet_name in sheet_names:
        alias, table_name = _allocate_table_names(sheet_name, table_map)

        try:
            con.execute(
                f"""
                CREATE TABLE {_quote_identifier(table_name)} AS
                SELECT * FROM st_read(
                    ?,
                    layer = ?,
                    open_options = ['HEADERS=FORCE', 'FIELD_TYPES=AUTO']
                )
            """,
                [file_path, sheet_name],
            )
            table_map[alias] = table_name
            row_count = con.execute(
                f"SELECT COUNT(*) FROM {_quote_identifier(table_name)}"
            ).fetchone()[0]
            logger.info(
                f"  Loaded sheet '{sheet_name}' -> table '{table_name}' ({row_count} rows)"
            )
        except Exception as e:
            logger.warning(f"  Failed to load sheet '{sheet_name}': {e}")


def _load_csv(
    con: duckdb.DuckDBPyConnection, file_path: str, table_map: dict[str, str]
) -> None:
    """Load a CSV file into a DuckDB table."""
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    alias, table_name = _allocate_table_names(base_name, table_map)

    try:
        con.execute(
            f"""
            CREATE TABLE {_quote_identifier(table_name)} AS
            SELECT * FROM read_csv_auto(?)
        """,
            [file_path],
        )
        table_map[alias] = table_name
        row_count = con.execute(
            f"SELECT COUNT(*) FROM {_quote_identifier(table_name)}"
        ).fetchone()[0]
        logger.info(
            f"  Loaded CSV '{base_name}' -> table '{table_name}' ({row_count} rows)"
        )
    except Exception as e:
        logger.warning(f"  Failed to load CSV '{base_name}': {e}")


def action_inspect(con: duckdb.DuckDBPyConnection, table_map: dict[str, str]) -> str:
    """Inspect the schema of all loaded tables."""
    output_parts = []

    for original_name, table_name in table_map.items():
        output_parts.append(f"\n{'=' * 60}")
        output_parts.append(f'Table: {original_name} (SQL name: "{table_name}")')
        output_parts.append(f"{'=' * 60}")

        # Get row count
        row_count = con.execute(
            f"SELECT COUNT(*) FROM {_quote_identifier(table_name)}"
        ).fetchone()[0]
        output_parts.append(f"Rows: {row_count}")

        # Get column info
        columns = con.execute(f"DESCRIBE {_quote_identifier(table_name)}").fetchall()
        output_parts.append(f"\nColumns ({len(columns)}):")
        output_parts.append(f"{'Name':<30} {'Type':<15} {'Nullable'}")
        output_parts.append(f"{'-' * 30} {'-' * 15} {'-' * 8}")
        for col in columns:
            col_name, col_type, nullable = col[0], col[1], col[2]
            output_parts.append(f"{col_name:<30} {col_type:<15} {nullable}")

        # Get non-null counts per column
        col_names = [col[0] for col in columns]
        non_null_parts = []
        for c in col_names:
            non_null_parts.append(f"COUNT({_quote_identifier(c)})")
        non_null_sql = (
            f"SELECT {', '.join(non_null_parts)} FROM {_quote_identifier(table_name)}"
        )
        try:
            non_null_counts = con.execute(non_null_sql).fetchone()
            output_parts.append("\nNon-null counts:")
            for i, c in enumerate(col_names):
                output_parts.append(f"  {c}: {non_null_counts[i]} / {row_count}")
        except Exception:
            pass

        # Sample data (first 5 rows)
        output_parts.append("\nSample data (first 5 rows):")
        try:
            sample = con.execute(
                f"SELECT * FROM {_quote_identifier(table_name)} LIMIT 5"
            ).fetchdf()
            output_parts.append(sample.to_string(index=False))
        except Exception:
            sample = con.execute(
                f"SELECT * FROM {_quote_identifier(table_name)} LIMIT 5"
            ).fetchall()
            header = [col[0] for col in columns]
            output_parts.append("  " + " | ".join(header))
            for row in sample:
                output_parts.append("  " + " | ".join(str(v) for v in row))

    result = "\n".join(output_parts)
    print(result)
    return result


def action_query(
    con: duckdb.DuckDBPyConnection,
    sql: str,
    table_map: dict[str, str],
    output_file: str | None = None,
) -> str:
    """Execute a SQL query and return/export results."""
    try:
        # Let DuckDB resolve identifiers without rewriting SQL literals, column
        # names, comments, or CTEs. Temporary views also work on read-only caches.
        for original_name, table_name in table_map.items():
            if original_name != table_name:
                con.execute(
                    f"CREATE OR REPLACE TEMP VIEW {_quote_identifier(original_name)} "
                    f"AS SELECT * FROM main.{_quote_identifier(table_name)}"
                )
        result = con.execute(sql)
        columns = [desc[0] for desc in result.description]
        rows = result.fetchall()
    except Exception as e:
        error_msg = f"SQL Error: {e}\n\nAvailable tables:\n"
        for orig, tbl in table_map.items():
            cols = con.execute(f"DESCRIBE {_quote_identifier(tbl)}").fetchall()
            col_names = [c[0] for c in cols]
            error_msg += f'  "{tbl}" ({orig}): {", ".join(col_names)}\n'
        print(error_msg)
        return error_msg

    # Format output
    if output_file:
        return _export_results(columns, rows, output_file)

    # Print as table
    return _format_table(columns, rows)


def _format_table(columns: list[str], rows: list[tuple]) -> str:
    """Format query results as a readable table."""
    if not rows:
        msg = "Query returned 0 rows."
        print(msg)
        return msg

    # Calculate column widths
    col_widths = [len(str(c)) for c in columns]
    for row in rows:
        for i, val in enumerate(row):
            col_widths[i] = max(col_widths[i], len(str(val)))

    # Cap column width
    max_width = 40
    col_widths = [min(w, max_width) for w in col_widths]

    # Build table
    parts = []
    header = " | ".join(str(c).ljust(col_widths[i]) for i, c in enumerate(columns))
    separator = "-+-".join("-" * col_widths[i] for i in range(len(columns)))
    parts.append(header)
    parts.append(separator)
    for row in rows:
        row_str = " | ".join(
            str(v)[:max_width].ljust(col_widths[i]) for i, v in enumerate(row)
        )
        parts.append(row_str)

    parts.append(f"\n({len(rows)} rows)")
    result = "\n".join(parts)
    print(result)
    return result


def _markdown_cell(value: object) -> str:
    """Keep table delimiters and line breaks inside a single Markdown cell."""
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\n", "<br>")
    )


def _export_results(columns: list[str], rows: list[tuple], output_file: str) -> str:
    """Export query results to a file (CSV, JSON, or Markdown)."""
    os.makedirs(os.path.dirname(output_file) or ".", exist_ok=True)
    ext = os.path.splitext(output_file)[1].lower()

    if ext == ".csv":
        import csv

        with open(output_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            writer.writerows(rows)

    elif ext == ".json":
        duplicates = [name for name, count in Counter(columns).items() if count > 1]
        if duplicates:
            raise ValueError(
                "JSON export requires unique column names; duplicates: "
                + ", ".join(repr(name) for name in duplicates)
                + ". Use SQL aliases (AS) to give each result column a distinct name."
            )
        records = []
        for row in rows:
            record = {}
            for i, col in enumerate(columns):
                val = row[i]
                # Handle non-JSON-serializable types
                if hasattr(val, "isoformat"):
                    val = val.isoformat()
                elif isinstance(val, (bytes, bytearray)):
                    val = val.hex()
                record[col] = val
            records.append(record)
        # Reject non-JSON numbers, including nested values, before touching the
        # destination so a failed export cannot overwrite an existing result.
        serialized = json.dumps(
            records, indent=2, ensure_ascii=False, default=str, allow_nan=False
        )
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(serialized)

    elif ext == ".md":
        with open(output_file, "w", encoding="utf-8") as f:
            # Header
            f.write("| " + " | ".join(_markdown_cell(c) for c in columns) + " |\n")
            f.write("| " + " | ".join("---" for _ in columns) + " |\n")
            # Rows
            for row in rows:
                f.write("| " + " | ".join(_markdown_cell(v) for v in row) + " |\n")
    else:
        msg = f"Unsupported output format: {ext}. Use .csv, .json, or .md"
        print(msg)
        return msg

    msg = f"Results exported to {output_file} ({len(rows)} rows)"
    print(msg)
    return msg


def action_summary(
    con: duckdb.DuckDBPyConnection,
    table_name: str,
    table_map: dict[str, str],
) -> str:
    """Generate statistical summary for a table."""
    # Resolve table name
    resolved = _quote_identifier(table_map.get(table_name, table_name))

    try:
        columns = con.execute(f"DESCRIBE {resolved}").fetchall()
    except Exception:
        available = ", ".join(f'"{t}" ({o})' for o, t in table_map.items())
        msg = f"Table '{table_name}' not found. Available tables: {available}"
        print(msg)
        return msg

    row_count = con.execute(f"SELECT COUNT(*) FROM {resolved}").fetchone()[0]

    output_parts = []
    output_parts.append(f"\nStatistical Summary: {table_name}")
    output_parts.append(f"Total rows: {row_count}")
    output_parts.append(f"{'=' * 70}")

    numeric_types = {
        "BIGINT",
        "INTEGER",
        "SMALLINT",
        "TINYINT",
        "DOUBLE",
        "FLOAT",
        "DECIMAL",
        "HUGEINT",
        "REAL",
        "NUMERIC",
    }

    for col in columns:
        col_name, col_type = col[0], col[1].upper()
        quoted_col = _quote_identifier(col_name)
        output_parts.append(f"\n--- {col_name} ({col[1]}) ---")

        # Check base type (strip parameterized parts)
        base_type = re.sub(r"\(.*\)", "", col_type).strip()

        if base_type in numeric_types:
            try:
                stats = con.execute(f"""
                    SELECT
                        COUNT({quoted_col}) as count,
                        AVG({quoted_col})::DOUBLE as mean,
                        STDDEV({quoted_col})::DOUBLE as std,
                        MIN({quoted_col}) as min,
                        QUANTILE_CONT({quoted_col}, 0.25) as q25,
                        MEDIAN({quoted_col}) as median,
                        QUANTILE_CONT({quoted_col}, 0.75) as q75,
                        MAX({quoted_col}) as max,
                        COUNT(*) - COUNT({quoted_col}) as null_count
                    FROM {resolved}
                """).fetchone()
                labels = [
                    "count",
                    "mean",
                    "std",
                    "min",
                    "25%",
                    "50%",
                    "75%",
                    "max",
                    "nulls",
                ]
                for label, val in zip(labels, stats):
                    if isinstance(val, float):
                        output_parts.append(f"  {label:<8}: {val:,.4f}")
                    else:
                        output_parts.append(f"  {label:<8}: {val}")
            except Exception as e:
                output_parts.append(f"  Error computing stats: {e}")
        else:
            try:
                stats = con.execute(f"""
                    SELECT
                        COUNT({quoted_col}) as count,
                        COUNT(DISTINCT {quoted_col}) as unique_count,
                        MODE({quoted_col}) as mode_val,
                        COUNT(*) - COUNT({quoted_col}) as null_count
                    FROM {resolved}
                """).fetchone()
                output_parts.append(f"  count   : {stats[0]}")
                output_parts.append(f"  unique  : {stats[1]}")
                output_parts.append(f"  top     : {stats[2]}")
                output_parts.append(f"  nulls   : {stats[3]}")

                # Show top 5 values
                top_vals = con.execute(f"""
                    SELECT {quoted_col}, COUNT(*) as freq
                    FROM {resolved}
                    WHERE {quoted_col} IS NOT NULL
                    GROUP BY {quoted_col}
                    ORDER BY freq DESC
                    LIMIT 5
                """).fetchall()
                if top_vals:
                    output_parts.append("  top values:")
                    for val, freq in top_vals:
                        pct = (freq / row_count * 100) if row_count > 0 else 0
                        output_parts.append(f"    {val}: {freq} ({pct:.1f}%)")
            except Exception as e:
                output_parts.append(f"  Error computing stats: {e}")

    result = "\n".join(output_parts)
    print(result)
    return result


def main():
    parser = argparse.ArgumentParser(description="Analyze Excel/CSV files using DuckDB")
    parser.add_argument(
        "--files",
        nargs="+",
        required=True,
        help="Paths to Excel (.xlsx/.xls) or CSV files",
    )
    parser.add_argument(
        "--action",
        required=True,
        choices=["inspect", "query", "summary"],
        help="Action to perform: inspect, query, or summary",
    )
    parser.add_argument(
        "--sql",
        type=str,
        default=None,
        help="SQL query to execute (required for 'query' action)",
    )
    parser.add_argument(
        "--table",
        type=str,
        default=None,
        help="Table name for summary (required for 'summary' action)",
    )
    parser.add_argument(
        "--output-file",
        type=str,
        default=None,
        help="Path to export results (CSV/JSON/MD)",
    )
    args = parser.parse_args()

    # Validate arguments
    if args.action == "query" and not args.sql:
        parser.error("--sql is required for 'query' action")
    if args.action == "summary" and not args.table:
        parser.error("--table is required for 'summary' action")

    try:
        with open_analysis_cache(args.files) as (con, table_map):
            if args.action == "inspect":
                action_inspect(con, table_map)
            elif args.action == "query":
                action_query(con, args.sql, table_map, args.output_file)
            elif args.action == "summary":
                action_summary(con, args.table, table_map)
    except ValueError as exc:
        sys.exit(f"Error: {exc}")


if __name__ == "__main__":
    main()
